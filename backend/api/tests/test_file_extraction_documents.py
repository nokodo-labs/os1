"""real document extraction through the auto loader routing.

builds genuine docx/pptx/xlsx/pdf bytes in-memory (no chat model involved)
and asserts the markitdown-backed extraction returns the document text.
"""

import io
import zipfile

import pytest
from openpyxl import Workbook
from pptx import Presentation

from api.settings import settings
from api.v1.service.files.text_contents import extraction as extraction_service
from api.v1.service.files.text_contents.extraction import (
	chunk_loaded_text,
	load_sdk_file_text,
)
from nokodo_ai.loaders import File as SDKFile


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_DOCX_PARAGRAPHS = (
	"Quarterly revenue grew nine percent on strong subscription renewals.",
	"Operating costs stayed flat while headcount increased by twelve people.",
	"The board approved the expanded research budget for the coming year.",
)

_PDF_TEXT = (
	"Utility scale solar deployment doubled last year as storage costs fell "
	"and grid operators cleared the interconnection backlog."
)


@pytest.fixture(autouse=True)
def _pin_auto_loader(monkeypatch: pytest.MonkeyPatch) -> None:
	monkeypatch.setattr(settings.assets.content_vectorization, "loader", "auto")


@pytest.fixture(autouse=True)
def _stub_embedding_capacity(monkeypatch: pytest.MonkeyPatch) -> None:
	"""pin chunk-size clamping; no embedding model is involved in these tests."""

	async def _capacity(session: object = None) -> int:
		return 8192

	monkeypatch.setattr(extraction_service, "embedding_token_capacity", _capacity)


def _docx_bytes(paragraphs: tuple[str, ...]) -> bytes:
	"""minimal OOXML word document readable by mammoth/markitdown."""
	body = "".join(f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>" for text in paragraphs)
	document = (
		'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
		"<w:document "
		'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
		f"<w:body>{body}</w:body></w:document>"
	)
	content_types = (
		'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
		'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
		'<Default Extension="rels" '
		'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
		'<Default Extension="xml" ContentType="application/xml"/>'
		'<Override PartName="/word/document.xml" ContentType="application/vnd.'
		'openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
		"</Types>"
	)
	rels = (
		'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
		"<Relationships "
		'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
		'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
		'officeDocument/2006/relationships/officeDocument" '
		'Target="word/document.xml"/>'
		"</Relationships>"
	)
	buffer = io.BytesIO()
	with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
		archive.writestr("[Content_Types].xml", content_types)
		archive.writestr("_rels/.rels", rels)
		archive.writestr("word/document.xml", document)
	return buffer.getvalue()


def _pptx_bytes(title: str, bullet: str) -> bytes:
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[1])
	slide.shapes.title.text = title
	slide.placeholders[1].text = bullet
	buffer = io.BytesIO()
	presentation.save(buffer)
	return buffer.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
	workbook = Workbook()
	sheet = workbook.active
	for row in rows:
		sheet.append(row)
	buffer = io.BytesIO()
	workbook.save(buffer)
	return buffer.getvalue()


def _pdf_bytes(text: str) -> bytes:
	"""minimal single-page pdf with a correct xref table."""
	content_stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
	objects = [
		b"<< /Type /Catalog /Pages 2 0 R >>",
		b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
		b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
		b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
		b"<< /Length "
		+ str(len(content_stream)).encode("ascii")
		+ b" >>\nstream\n"
		+ content_stream
		+ b"\nendstream",
		b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
	]
	out = bytearray(b"%PDF-1.4\n")
	offsets: list[int] = []
	for index, body in enumerate(objects, start=1):
		offsets.append(len(out))
		out += f"{index} 0 obj\n".encode("ascii") + body + b"\nendobj\n"
	xref_at = len(out)
	out += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
	out += b"0000000000 65535 f \n"
	for offset in offsets:
		out += f"{offset:010d} 00000 n \n".encode("ascii")
	out += (
		f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
		f"startxref\n{xref_at}\n%%EOF\n"
	).encode("ascii")
	return bytes(out)


async def test_docx_extracts_paragraph_text() -> None:
	loaded = await load_sdk_file_text(
		SDKFile(
			data=_docx_bytes(_DOCX_PARAGRAPHS),
			filename="report.docx",
			mime_type=DOCX_MIME,
		)
	)

	assert loaded.status == "loaded"
	assert "markitdown" in loaded.source
	for paragraph in _DOCX_PARAGRAPHS:
		assert paragraph in loaded.content


async def test_docx_text_chunks_with_offsets() -> None:
	loaded = await load_sdk_file_text(
		SDKFile(
			data=_docx_bytes(_DOCX_PARAGRAPHS),
			filename="report.docx",
			mime_type=DOCX_MIME,
		)
	)
	chunks = await chunk_loaded_text(loaded)

	assert chunks
	first = chunks[0]
	char_start = first.metadata["char_start"]
	char_end = first.metadata["char_end"]
	assert isinstance(char_start, int)
	assert isinstance(char_end, int)
	assert loaded.content[char_start:char_end].strip() == first.text.strip()


async def test_pptx_extracts_slide_text() -> None:
	title = "Expansion Roadmap For Northern Markets"
	bullet = "Open two regional distribution hubs before the winter season."
	loaded = await load_sdk_file_text(
		SDKFile(
			data=_pptx_bytes(title, bullet),
			filename="roadmap.pptx",
			mime_type=PPTX_MIME,
		)
	)

	assert loaded.status == "loaded"
	assert title in loaded.content
	assert bullet in loaded.content


async def test_xlsx_extracts_cell_text() -> None:
	rows = [
		["region", "revenue", "note"],
		["north", "125000", "steady growth across all quarters"],
		["south", "98000", "seasonal demand dipped in summer"],
	]
	loaded = await load_sdk_file_text(
		SDKFile(data=_xlsx_bytes(rows), filename="revenue.xlsx", mime_type=XLSX_MIME)
	)

	assert loaded.status == "loaded"
	assert "steady growth across all quarters" in loaded.content
	assert "98000" in loaded.content


async def test_pdf_extracts_page_text() -> None:
	loaded = await load_sdk_file_text(
		SDKFile(
			data=_pdf_bytes(_PDF_TEXT),
			filename="grid-report.pdf",
			mime_type="application/pdf",
		)
	)

	assert loaded.status == "loaded"
	assert "solar deployment doubled" in loaded.content
